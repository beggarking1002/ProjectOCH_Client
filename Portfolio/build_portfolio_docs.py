from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path(__file__).resolve().parent
FONT = "Malgun Gothic"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(89, 99, 110)


def set_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if italic is not None:
        run.italic = italic


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.2

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("PROJECT OCH | PORTFOLIO DRAFT")
    set_font(run, size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("ProjectOCH - Portfolio Draft")
    set_font(run, size=8, color=MUTED)
    return doc


def title_block(doc, label, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(label.upper())
    set_font(r, size=10, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    set_font(r, size=25, bold=True, color=RGBColor(11, 37, 69))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run(subtitle)
    set_font(r, size=12, color=MUTED)


def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def paragraph(doc, text, lead=None):
    p = doc.add_paragraph()
    if lead:
        r = p.add_run(lead)
        set_font(r, bold=True, color=DARK_BLUE)
    r = p.add_run(text)
    set_font(r)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_font(r)


def numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_font(r)


def callout(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    ppr = p._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F4F6F9")
    ppr.append(shading)
    r = p.add_run(f"{label}  ")
    set_font(r, bold=True, color=DARK_BLUE)
    r = p.add_run(text)
    set_font(r)


def build_planning_portfolio():
    doc = setup_document()
    title_block(
        doc,
        "Planning Portfolio",
        "ProjectOCH 기획 포트폴리오 초안",
        "필드에서 전투까지 연결한 서버 권위형 육각 타일 RPG의 경험 설계",
    )
    callout(
        doc,
        "포트폴리오 포지셔닝",
        "ProjectOCH는 '서버가 판정한 복잡한 전투를 플레이어가 행동 전·중·후에 이해하도록 만드는 게임'으로 제시한다. "
        "필드 이동과 전투 초대, 씬 전환, 육각 타일 전투, 결과 복귀가 하나의 플레이 루프로 구현되어 있어, 전투 기능 나열보다 경험의 연결을 보여 주기에 적합하다.",
    )

    heading(doc, "1. 프로젝트 한눈에 보기")
    paragraph(doc, "ProjectOCH는 Unity로 제작한 온라인 RPG 클라이언트다. 타이틀에서 게임 입장을 요청하고, 필드에서 다른 플레이어의 이동과 전투 초대를 처리한 뒤, "
                   "전투 씬에서 서버 권위형 육각 타일 턴제 전투를 진행하고 결과 확인 후 필드로 돌아온다. 서버는 이동·명중·회피·방어·반격·상태·타일 상태의 최종 결과를 판정하며, "
                   "클라이언트는 그 결과를 입력 경험과 시각 피드백으로 번역한다.")
    bullets(doc, [
        "플레이 흐름: 타이틀 → 필드 입장/이동 → 전투 초대 → 전투 씬 → 결과 확인 → 필드 복귀",
        "공간: 필드는 Unity Tilemap 기반, 전투는 axial 좌표 기반 육각 타일",
        "전투 요소: 이동, 슬롯형 스킬, 상태 이상, HP/Armor/자원, 타일 오버레이, 장비 타일, 반격, ZOC",
        "데이터 현황: 클래스 매핑·Pawn 템플릿은 8종, 완성된 BattleSkill 데이터 세트는 Beige Ice·Beige Fire·Suen Axe·Zillian Longbow·Alen Spear 5종",
        "콘텐츠 파이프라인: CSV GameData와 Addressables를 통해 수치·표현·에셋을 연결",
    ])

    heading(doc, "2. 기획 목표와 핵심 문제")
    paragraph(doc, "이 프로젝트의 핵심 과제는 서버가 관리하는 복잡한 규칙을 클라이언트에서 임의로 판정하지 않으면서도, 플레이어가 납득 가능한 선택을 하게 만드는 것이다. "
                   "따라서 기획은 '행동 전의 예측성', '행동 중의 가독성', '행동 후의 인과성'을 중심으로 구성했다.")
    bullets(doc, [
        "행동 전: 사거리, 타겟 타입, 빈 타일 조건, 타일 오버레이 조건, 실제 효과 범위를 한 화면에서 파악한다.",
        "이동 전: 현재 위치에서 이탈할 때 반응할 적의 ZOC를 보고 위험을 선택한다.",
        "행동 중: 일반 타격·회피·방어·반격·재반격이 섞여도 각 타격의 주체와 대상, 순서를 놓치지 않는다.",
        "행동 후: HP·Armor·COLD/HEAT/MORALE·상태·도끼/불 타일 등 서버 delta의 변화를 단계적으로 확인한다.",
        "전투 밖: 타이틀과 필드의 네트워크 흐름을 끊지 않고 초대와 결과 복귀까지 자연스럽게 연결한다.",
    ])

    heading(doc, "3. 전체 플레이 경험 설계")
    heading(doc, "3.1 필드에서 전투로 넘어가는 온라인 경험", level=2)
    paragraph(doc, "타이틀의 시작 버튼은 C_ENTER_GAME을 보내고, 성공 응답 뒤 전환 오버레이를 거쳐 FieldScene을 연다. 필드에서는 고정소수점 좌표 패킷으로 이동을 요청하고, "
                   "서버가 내려준 이동 시간에 맞춰 다른 플레이어의 이동을 재생한다. 전투 초대 요청·수신·결과 UI와 S_ENTER_BATTLE 연동을 두어, 전투를 독립 샌드박스가 아닌 필드 활동의 결과로 다룬다.")
    bullets(doc, [
        "필드 클릭은 Tilemap의 Ground/Block 정보를 이용해 로컬에서 1차 이동 가능 여부를 확인하고, 최종 이동은 서버에 요청",
        "FieldPositionCodec이 월드 좌표를 100배 고정소수점 Vec2로 변환해 네트워크 표현을 일관화",
        "전투 초대 팝업은 필드 입력을 차단해 의도치 않은 이동과 UI 조작이 겹치지 않게 설계",
        "씬 전환 오버레이와 Addressables 로더로 필드/전투 리소스의 진입·해제를 관리",
    ])

    heading(doc, "3.2 타겟 프리뷰를 규칙의 언어로 사용", level=2)
    paragraph(doc, "스킬의 유효 타일과 실제 영향 타일을 분리해 표현했다. 유효 타일은 선택 가능 범위, 영향 타일은 결과 프리뷰로 표시하여, "
                   "플레이어가 '어디를 고를 수 있는가'와 '고르면 무엇이 일어나는가'를 동시에 이해하도록 했다.")
    bullets(doc, [
        "RADIUS_1: 중심과 인접 6타일을 동시에 제시",
        "LINE_3: 시전자와 선택 방향을 기준으로 일직선 결과를 제시",
        "Fire Wall: 시작 타일 선택 후 인접 6방향 중 두 번째 타일을 선택하는 2단계 입력",
        "Teleport: 불타는 빈 타일만 유효 후보로 제한",
        "SELF/ALLY_OR_SELF/PICKUP_TILE 등 TargetType에 따라 입력 대상과 표시를 다르게 처리",
        "Suen Cleaner: 자기 자신 또는 인접 범위를 다시 선택해 발동하는 자기 중심 광역 입력",
    ])

    heading(doc, "3.3 캐릭터 정체성을 의사결정으로 연결", level=2)
    paragraph(doc, "캐릭터는 단순 수치 차이가 아니라 타일, 자원, 장비 상태와 결합한 선택 경험으로 설계했다. Beige Fire는 HEAT를 쌓아 강한 행동을 사용하면서 과열 위험을 관리하고, "
                   "Beige Ice는 COLD 자원과 범위형 스킬을 사용한다. Suen Axe는 도끼 보유 여부에 따라 슬롯의 명칭·아이콘·행동 성격이 바뀌며, Alen Spear는 전방 2칸 ZOC로 위치 압박을 만든다.")
    bullets(doc, [
        "Beige Fire: Fireball, Explosion, Fire Wall, FIRE 타일 텔레포트, HEAT 관리",
        "Beige Ice: Ice Bolt, Shield, Hail, Storm Center 등 냉기 기반의 원거리 선택 경험",
        "Suen Axe: AxeOn/AxeOff 상태별 슬롯 표현, 도끼 투척 후 AXE 타일 표시, 소유 도끼만 회수 가능",
        "Zillian Longbow: 원거리 공격/버프 중심의 별도 BattleSkill 세트",
        "Alen Spear: 일반 근접보다 넓은 전방 부채꼴 2칸 ZOC와 반응 공격",
    ])

    heading(doc, "3.4 결과를 한 번에 정산하지 않는 전투 피드백", level=2)
    paragraph(doc, "반격 체인은 서버가 전달한 BattleActionLog의 순서를 그대로 재생한다. 각 타격마다 0.5초 간격으로 스킬 스프라이트, "
                   "근접 돌진, 회피, 데미지 숫자, HP/실드 바를 갱신해 '누가 누구를 때렸는지'를 시각적으로 분해했다.")
    bullets(doc, [
        "회피: 공격자 반대 방향으로 두 단계 점프 후 복귀",
        "근접 공격 및 반격: 대상 방향으로 짧게 돌진 후 복귀",
        "ZOC 반응: 이동 완료 후 반응 공격과 이어지는 반격을 서버 로그 순서로 재생",
        "회피는 evaded일 때만 적용하며, Armor로 막힌 경우는 제자리에서 결과만 갱신",
        "전투 연출 중 입력을 잠가 중복 클릭과 클라이언트 상태 경합을 방지",
    ])

    heading(doc, "3.5 ZOC를 이동 위험도 UI로 번역", level=2)
    paragraph(doc, "ZOC는 목적지가 아니라 '현재 위치에서 이탈할 때'의 반응 조건으로 해석했다. 이동 가능 타일에 마우스를 올리면, "
                   "현재 위치가 적의 LEAVE_ZONE ZOC 안에 있고 적의 반응 횟수가 남았을 때만 해당 적과 전방 부채꼴 범위를 표시한다.")
    bullets(doc, [
        "Suen Axe: 전방 부채꼴 1칸, LEAVE_ZONE 반응 1회",
        "Alen Spear: 전방 부채꼴 2칸, LEAVE_ZONE 반응 1회",
        "반응 가능 적은 빨강, 해당 적의 ZOC 타일은 노랑으로 표시",
        "범위·전방 폭·반응 횟수·반응 슬롯·트리거는 BattleZoc 데이터 테이블로 관리",
    ])

    heading(doc, "4. 데이터 구조를 활용한 콘텐츠 운영 관점")
    paragraph(doc, "GameData는 ClassKey, PawnTemplate, BattleSkill, BattleSkillEffect, BattleSkillEffectParam, BattleSkillView, BattleZoc, DisplayText, EnumDef로 나뉜다. "
                   "즉 캐릭터 기본 능력, 스킬 타겟/범위, 효과 조합, 아이콘/애니메이션, ZOC, 다국어 텍스트와 열거형 설명을 분리해 두었다. 신규 콘텐츠는 코드 수정보다 데이터와 에셋 매핑을 먼저 검토하는 방식으로 확장할 수 있다.")
    bullets(doc, [
        "PawnTemplate: 역할, 기본 능력치, 이동력, 자원, 반격 슬롯 등 Pawn의 전투 베이스 정의",
        "BattleSkill/BattleSkillEffect/Param: 스킬 선택 규칙과 효과 조합을 분리",
        "BattleSkillView: 스킬 키와 Animator/VFX/SFX/Icon 키를 연결",
        "DisplayText/EnumDef: UI 표기와 설명의 데이터화 기반",
        "Addressables: GameData, 맵, Pawn 프리팹, 타일, UI, 스킬 아이콘을 주소 키로 관리",
    ])

    heading(doc, "5. 포트폴리오 경쟁력 평가")
    paragraph(doc, "기획 포트폴리오로서의 기본 경쟁력은 있다. 가장 큰 강점은 전투 규칙을 설명으로 끝내지 않고, 타겟 프리뷰·ZOC 경고·상태/UI·전투 로그 연출처럼 플레이어가 실제로 접하는 접점까지 구현했다는 점이다. "
                   "또한 필드와 전투를 연결한 온라인 흐름이 있어, 단일 전투 화면보다 게임 전체를 바라본 기획으로 제시할 수 있다.")
    bullets(doc, [
        "경쟁력 있는 지점: 서버 권위형 제약, 육각 타일 공간, 자원/상태/타일 상호작용을 사용성 문제로 풀어낸 사례",
        "경쟁력 있는 지점: 캐릭터 정체성을 데이터, 입력 규칙, 아이콘, 애니메이션까지 일관되게 연결한 사례",
        "보완할 지점: 클래스 8종 중 스킬 데이터가 완결된 5종이라는 현재 범위를 명확히 밝히고, 확장 계획을 제시할 것",
        "보완할 지점: 본인 역할과 의사결정 근거를 기능별로 구분하고, 실제 플레이테스트 수치가 생기면 결과를 추가할 것",
        "보완할 지점: 제출물에는 20~30초 영상으로 필드 초대 → 전투 진입 → 범위 선택 → 반격/ZOC → 결과 복귀를 보여 줄 것",
    ])

    heading(doc, "6. 제출용 구성 제안")
    paragraph(doc, "문서만으로는 상호작용의 설득력이 제한되므로, 아래 구성을 함께 제출하는 것이 좋다. 이는 현재 코드에 있는 기능을 근거로 한 제안이며, 아직 측정하지 않은 사용자 지표는 임의로 기재하지 않는다.")
    numbered(doc, [
        "1장: 문제 정의와 전체 플레이 루프. 타이틀·필드·전투·복귀 흐름을 한 장의 다이어그램으로 제시",
        "2장: Fire Wall 2단계 타게팅 또는 텔레포트의 조건부 타겟 프리뷰를 전/후 스크린샷으로 제시",
        "3장: ZOC 경고와 반격 체인을 GIF/영상으로 제시하고, 왜 목적지가 아닌 현재 위치를 기준으로 했는지 설명",
        "4장: CSV 한 행이 스킬 아이콘·타겟 범위·전투 규칙으로 연결되는 콘텐츠 제작 흐름을 제시",
        "마무리: 플레이 테스트를 진행한 뒤에는 오선택률·규칙 이해도·연출 인지 여부 같은 실제 결과만 추가",
    ])

    heading(doc, "7. 구현 근거가 되는 알고리즘과 시스템 설계")
    paragraph(doc, "ProjectOCH의 강점은 특정 유명 알고리즘을 나열하는 데 있지 않다. 온라인 턴제 전투의 실제 문제에 맞춰 좌표 계산, 범위 확장, 서버 상태 동기화, "
                   "이벤트 순서 재생을 결합한 데 있다. 기획 포트폴리오에서는 이를 플레이어가 예측하고 납득하는 경험을 만든 구현 근거로 제시한다.")
    bullets(doc, [
        "육각 타일 거리 계산: axial 좌표를 cube 거리로 환산해 이동 가능 범위, 사거리, 스킬 범위 검사를 같은 규칙으로 처리",
        "ZOC 전방 부채꼴 확장: Pawn의 바라보는 방향에서 전방/좌전방/우전방으로 거리별 확장을 반복하고, HashSet으로 중복 타일을 제거",
        "창병 ZOC: 거리 1의 3칸과 거리 2의 5칸으로 넓어지는 총 8칸의 전방 부채꼴을 만들어 위치 압박을 표현",
        "서버 delta 동기화: pawn_deltas와 tile_deltas를 단일 반영 경로로 적용해 HP·Armor·자원·상태·좌표·밀쳐내기를 같은 권위 데이터로 처리",
        "로그 순서 재생: BattleActionLog 순서를 보존해 일반 공격, 회피, 반격, ZOC 공격의 원인과 UI 변화를 타격 단위로 분해",
        "메인 스레드 패킷 큐: 비동기 소켓 수신을 Unity 오브젝트 접근과 분리해 화면 갱신의 스레드 안전성을 확보",
    ])
    doc.save(OUT_DIR / "ProjectOCH_기획_포트폴리오_초안.docx")


def build_development_portfolio():
    doc = setup_document()
    title_block(
        doc,
        "Development Portfolio",
        "ProjectOCH 기술소개서",
        "Unity 온라인 클라이언트의 필드·전투·데이터·네트워크 아키텍처",
    )
    callout(
        doc,
        "기술 요약",
        "ProjectOCH는 TCP/Protocol Buffers 기반 서버와 통신하는 Unity 클라이언트다. 네트워크 수신은 메인 스레드 큐로 전환하고, "
        "필드와 전투는 서버의 패킷·delta·action log를 기준으로 상태와 표현을 동기화한다.",
    )

    heading(doc, "1. 코드베이스 기준 기술 스택과 범위")
    bullets(doc, [
        "Runtime: Unity / C# / uGUI / TextMeshPro / Tilemap / Animator",
        "Networking: TCP SocketAsyncEventArgs, 자체 Session·RecvBuffer·PacketManager, Protocol Buffers 생성 코드",
        "State model: 서버 권위형 필드 및 전투 상태, battle_state_version, pawn_deltas, tile_deltas, BattleActionLog",
        "Content: CSV GameData, Unity Addressables, Pawn 프리팹·타일·UI·스킬 아이콘",
        "Gameplay: axial 육각 좌표, 턴제 스킬, 상태/자원/Armor, 타일 오버레이, 장비 타일, 반격, ZOC",
        "프로젝트 규모(작성 코드 기준): BattleObjectManager 약 1,865줄, BattleUIController 약 1,826줄, BattleGameDataRepository 약 903줄의 핵심 모듈",
    ])

    heading(doc, "2. 런타임 구성과 네트워크 경계")
    paragraph(doc, "GameRoot는 씬을 넘겨 유지되는 앱 루트다. AppServices를 통해 NetworkService를 초기화하고 매 프레임 Tick을 호출하며, "
                   "Addressables에서 Maplestory Light 폰트를 로드해 씬의 uGUI·TMP·월드 TextMesh에 적용한다. GameServerConnection은 씬 컴포넌트가 사용할 수 있는 편의 façade로 동작한다.")
    bullets(doc, [
        "Connector와 Session: 비동기 TCP 연결, 송신 큐/보류 버퍼, 수신 버퍼, 4바이트 패킷 헤더(size + packetId) 처리",
        "GameServerSession: 패킷을 PacketManager로 넘기고 연결·해제 이벤트를 NetworkService에 전달",
        "PacketHandler: 네트워크 콜백에서 받은 S_* 패킷을 lock 기반 메인 스레드 작업 큐에 넣고 Flush로 Unity 객체 접근을 직렬화",
        "NetworkService: 로그인, 게임 입장, 필드 spawn/despawn/move, 전투 초대, 전투 이동/스킬/턴 종료/결과 ACK의 송수신과 상태 이벤트를 집약",
        "생성 코드와 작성 코드 분리: Assets/Scripts/Packet/Generated의 Protocol·Struct·Enum·PacketManager는 proto 재생성 산출물이며, 게임 표현 로직은 별도 모듈에 둔다.",
    ])

    heading(doc, "3. 필드·씬 전환 구현")
    paragraph(doc, "TitleSceneFlow는 시작 버튼을 C_ENTER_GAME에 연결하고, 성공 패킷을 받으면 전환 오버레이 후 FieldScene을 로드한다. "
                   "FieldSceneAddressableLoader는 필드 맵과 월드 맵을 Addressables로 로드하며, FieldObjectManager는 서버 spawn/despawn/move 패킷에 맞춰 Pawn을 생성·갱신한다. "
                   "FieldPawnController는 입력을 월드 좌표로 변환하고 Tilemap 기반 보행 가능 여부를 1차 검증한 뒤 C_MOVE를 보낸다.")
    bullets(doc, [
        "FieldPositionCodec: float 월드 좌표 ↔ 100배 고정소수점 Vec2 변환",
        "FieldMapWalkArea / FieldMapAxialCoordinates: Grid·Ground Tilemap·Block/Prop Tilemap 기반의 보행 및 좌표 질의",
        "FieldBattleInviteUI: 전투 초대 요청/수신/결과를 UI로 처리하고 입력을 일시 차단",
        "BattleSceneFlow: S_ENTER_BATTLE 수신 시 BattleScene 진입, 전투 결과 ACK 뒤 FieldScene 복귀",
        "SceneTransitionOverlay: Addressables 프리팹을 비동기 생성하고 전환 중 중복 생성/해제를 방지",
    ])

    heading(doc, "4. 전투 아키텍처와 동기화")
    paragraph(doc, "전투는 클라이언트 예측으로 결과를 확정하지 않는다. S_ENTER_BATTLE로 초기 Pawn·타일 상태를 구성하고, S_BATTLE_MOVE·S_BATTLE_SKILL·S_BATTLE_END_TURN의 pawn_deltas와 tile_deltas로 상태를 갱신한다. "
                   "BattleObjectManager는 현재 전투와 state version을 확인하고, 전투 연출 중에는 입력을 잠가 수신 순서와 화면 상태가 충돌하지 않게 제어한다.")
    bullets(doc, [
        "BattleMapGrid: axial 좌표/월드 좌표 변환, 타일 상태, 이동 가능 여부, FIRE 오버레이 및 AXE 장비 타일 표현",
        "BattlePawn: HP·Armor·상태·자원·행동 플래그·is_action_blocked·ZOC 사용 횟수를 서버 상태에서 반영",
        "BattlePawn 계층: BattlePawn → Beige → BeigeIce/BeigeFire, BattlePawn → Suen → SuenAxe로 공통과 고유 표현을 분리",
        "BattleTargetPreview: 이동, 스킬 사거리, RADIUS_1, LINE_3, 2단계 Fire Wall, 조건부 Teleport, ZOC 위험을 화면에 표시",
        "BattleUIController: 턴 정보, 슬롯, 선택/대상 상세, HP/Armor, COLD/HEAT/MORALE, 상태 스택, 전투 숫자를 표시",
    ])

    heading(doc, "5. CSV GameData와 Addressables 콘텐츠 파이프라인")
    paragraph(doc, "BattleGameDataRepository는 Addressables에서 CSV를 비동기로 로드하고, 에디터에서는 Assets/GameData 경로를 fallback으로 사용한다. "
                   "파싱 결과는 PawnClass·ClassKey·스킬 키·슬롯 기준 딕셔너리로 보관해 런타임 조회 비용과 콘텐츠 의존성을 낮춘다.")
    bullets(doc, [
        "ClassKey.csv / PawnTemplate.csv: 8개 Pawn 클래스의 역할, 기본 능력치, 이동력, 자원, 반격 슬롯을 매핑",
        "BattleSkill.csv: SkillCategory, ActionSlot, AP, 최소/최대 사거리, TargetType, TargetShape, RequiredOverlayType, 효과 그룹 정의",
        "BattleSkillEffect.csv / Param.csv: 효과 인스턴스와 파라미터를 분리해 복합 효과를 순서대로 조합",
        "BattleSkillView.csv: 스킬 키를 Animator trigger, VFX, SFX, 아이콘 주소 키와 연결",
        "BattleZoc.csv: PawnClass별 범위, 전방 폭, 반응 한도, 반응 스킬, 트리거를 정의",
        "Addressables: GameData 외에 맵, Pawn 프리팹, 타일, UI, 스킬 아이콘을 개별 키로 로드·해제",
    ])

    heading(doc, "6. 구현 사례: 서버 로그를 읽을 수 있는 전투 연출로 변환")
    paragraph(doc, "일반 스킬 하나에는 일반 타격, 회피, 가드, 반격, 재반격, ZOC 반응이 연쇄될 수 있다. 최종 HP/Armor만 즉시 적용하면 원인과 순서가 사라지므로, "
                   "BattleActionLog의 패킷 순서를 보존한 Coroutine 시퀀스로 각 타격을 0.5초 단위로 재생한다. UI 수치와 게이지도 각 로그 시점에 함께 갱신한다.")
    bullets(doc, [
        "각 로그: 공격자 스킬 표현 → 회피/근접 돌진 → 대상의 HP·Armor·전투 숫자 갱신",
        "is_evaded: 타격 대상이 뒤로 폴짝 이동한 뒤 앞으로 복귀. Armor 방어에는 회피 모션을 적용하지 않음",
        "is_counter: 반격 로그도 패킷 순서대로 재생하므로 재반격 체인을 자연스럽게 연결",
        "ZOC: 이동 후 내려온 반응 로그를 같은 시퀀스에 넣어 창병/근접 Pawn의 돌진으로 표현",
        "연출 잠금: 공격·반격·UI 갱신이 진행되는 동안 타겟 입력과 슬롯 조작을 차단",
    ])

    heading(doc, "7. 구현 사례: 상태 기반 입력과 캐릭터 특화 표현")
    paragraph(doc, "공용 컨트롤러에 모든 캐릭터 규칙을 넣지 않고, 공통 상태 반영은 BattlePawn에 두며 캐릭터 고유 표현은 BeigeIce, BeigeFire, SuenAxe 같은 파생 클래스에 둔다. "
                   "이 구조는 새 캐릭터가 추가될 때 공용 전투 흐름을 훼손하지 않고 Animator·아이콘·고유 타게팅만 확장하도록 돕는다.")
    bullets(doc, [
        "BeigeFire: FIRE 오버레이 텔레포트, RADIUS_1 Explosion, LINE_3 Fire Wall의 2단계 방향 타겟팅",
        "SuenAxe: SUEN_AXE_AXE_OFF 상태를 Animator와 AxeOn/AxeOff 아이콘 매핑에 연결하고, AXE 장비 타일 소유자를 확인해 회수 후보를 제한",
        "일반 스킬·궁극기·보조 행동: remaining AP만으로 전체를 막지 않고 used_normal_skill_this_turn / used_ultimate / used_sub_action_this_turn을 각 슬롯 규칙에 반영",
        "is_action_blocked: 이동과 슬롯 2~7을 비활성화하고 턴 종료만 남겨, 상태 키가 아니라 서버가 보낸 행동 제한 플래그를 기준으로 제어",
        "리소스·상태: COLD, HEAT, MORALE의 현재/최대값과 DIZZY 등 status stacks / remaining owner turns를 패널에 표시",
    ])

    heading(doc, "8. 구현 사례: 전투 규칙을 위한 알고리즘")
    paragraph(doc, "이 프로젝트의 알고리즘은 게임 규칙과 표현의 일관성을 지키기 위한 실무형 구현이다. 모든 계산은 서버 판정을 대체하지 않으며, "
                   "입력 가능 범위와 결과 예측을 제공하고 서버가 내려 준 최종 상태를 읽기 쉬운 연출로 전환하는 데 사용한다.")
    bullets(doc, [
        "Axial cube distance: abs(dq), abs(dr), abs(ds)의 합을 2로 나누는 육각 거리로 이동·사거리·타겟 검사의 기준을 통일",
        "방향 기반 BFS ZOC: 전방 폭 3을 기준으로 frontier를 거리별 확장하고, 맵 경계 검사와 HashSet 중복 제거로 전방 부채꼴 타일을 생성",
        "현재 위치 기준 LEAVE_ZONE 판정: 마우스가 올린 목적지는 이동 가능 여부만 확인하고, 반응 가능성은 이동 시작 Pawn의 현재 axial이 적 ZOC에 포함되는지로 판정",
        "서버 delta 위치 보간: BattlePawnDelta.axial 변화는 MoveToAxial로 연결해 텔레포트·밀쳐내기·충돌 뒤 제자리 유지도 하나의 좌표 반영 경로로 처리",
        "순서 보존 액션 시퀀서: BattleActionLog 배열을 재정렬하지 않고 Coroutine으로 순차 재생하여 반격/재반격 및 Sentinel ZOC 개입을 타격별 UI 갱신과 결합",
        "메인 스레드 작업 큐: lock으로 보호한 큐에 패킷 핸들러 작업을 넣고 Tick/Flush에서 실행해 Unity API 접근을 네트워크 콜백과 분리",
    ])

    heading(doc, "9. 품질 확인과 기술 부채")
    paragraph(doc, "이 프로젝트는 서버와 함께 동작하는 게임 클라이언트이므로, 컴파일 검증 외에 실제 패킷의 상태값과 로그 순서를 확인하는 통합 검증이 중요하다. "
                   "현재 코드베이스에서는 전용 자동화 테스트 어셈블리보다 수동/통합 검증 중심의 형태가 확인된다. 따라서 아래 항목은 현재의 강점과 다음 개선 과제를 구분해 제시한다.")
    bullets(doc, [
        "빌드 검증: dotnet build Assembly-CSharp.csproj --no-restore",
        "상태 검증: pawn delta의 HP, armor, resource, status, action flag 반영",
        "표현 검증: 타겟 프리뷰 레이어, Pawn sorting, 아이콘 로드, Animator 상태 전환",
        "연출 검증: 일반 타격, 회피, 반격, ZOC 반응의 순서 및 입력 잠금",
        "기술 부채: BattleObjectManager와 BattleUIController가 큰 조정자 역할을 맡고 있어, 입력·패킷 적용·연출·UI를 더 작은 책임으로 분리할 여지가 있음",
        "기술 부채: 일부 필드/전투 기본 Addressables 주소가 Beige Fire에 고정되어 있어, 계정/파티 기반 Pawn 선택 데이터로 치환할 필요가 있음",
        "개선 제안: 고정된 S_BATTLE_* fixture로 타겟 조건, action log 순서, delta 적용을 검증하는 편집 모드/플레이 모드 테스트 추가",
    ])

    heading(doc, "10. 개발 포트폴리오로서의 핵심 어필")
    paragraph(doc, "이 코드베이스가 보여 주는 핵심 역량은 단순 Unity 화면 구현이 아니라, 비동기 네트워크·서버 권위 상태·데이터 테이블·Addressables·Animator/UI를 한 게임 루프로 연결한 점이다. "
                   "포트폴리오에서는 생성된 proto 코드의 양보다, 그 계약을 안전하게 수신하고 플레이어가 읽을 수 있는 피드백으로 만드는 작성 코드의 판단을 중심으로 설명해야 한다.")
    bullets(doc, [
        "강점: Socket 수신을 메인 스레드 큐로 넘기고, 씬 전환까지 지속되는 네트워크 서비스로 라이프사이클을 구성",
        "강점: 서버 결과의 delta와 action log를 분리해 상태 일관성과 전투 가독성을 함께 확보",
        "강점: CSV/Addressables/프로토콜/프리팹을 분리해 콘텐츠 확장 지점을 코드에서 분명히 함",
        "강점: 공용 Pawn 흐름과 캐릭터별 표현을 계층화해 전투 규칙의 공용화와 개성 표현을 양립",
        "다음 증빙: 패킷 캡처, Unity Profiler 수치, 멀티 클라이언트 필드/초대 영상, 자동화 테스트 결과를 추가하면 설득력이 더 높아짐",
    ])
    doc.save(OUT_DIR / "ProjectOCH_개발_기술소개서.docx")


if __name__ == "__main__":
    build_planning_portfolio()
    build_development_portfolio()
